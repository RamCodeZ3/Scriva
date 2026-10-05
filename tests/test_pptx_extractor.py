from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from domain.entities.source import FileKind, SourceType, classify_source
from infrastructure.extractors.file_extractor_adapter import (
    FileExtractorAdapter,
)
from pptx import Presentation
from pptx.util import Inches


class PptxExtractorTest(unittest.IsolatedAsyncioTestCase):
    def test_classifies_pptx_as_document_file(self) -> None:
        source_type, file_kind = classify_source("presentation.pptx")

        self.assertEqual(source_type, SourceType.FILE)
        self.assertEqual(file_kind, FileKind.DOCUMENT)

    async def test_extracts_slides_and_shapes_in_reading_order(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        lower = slide.shapes.add_textbox(
            Inches(1), Inches(3), Inches(4), Inches(1)
        )
        lower.text = "Second idea"
        upper = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(4), Inches(1)
        )
        upper.text_frame.paragraphs[0].text = "First idea"
        paragraph = upper.text_frame.add_paragraph()
        paragraph.text = "First detail"

        table_slide = presentation.slides.add_slide(
            presentation.slide_layouts[6]
        )
        table = table_slide.shapes.add_table(
            2, 2, Inches(1), Inches(1), Inches(5), Inches(2)
        ).table
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Alpha"
        table.cell(1, 1).text = "10"

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ordered.pptx"
            presentation.save(path)
            result = await FileExtractorAdapter().extract(str(path))

        self.assertEqual(
            result,
            "[Slide 1]\nFirst idea\nFirst detail\nSecond idea\n\n"
            "[Slide 2]\nName\tValue\nAlpha\t10",
        )


if __name__ == "__main__":
    unittest.main()
